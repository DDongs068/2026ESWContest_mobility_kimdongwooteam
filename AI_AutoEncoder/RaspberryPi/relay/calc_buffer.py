#!/usr/bin/env python3
import argparse
import os
import socket
import subprocess
import time

import uno
from com.sun.star.beans import PropertyValue

HERE = os.path.dirname(os.path.abspath(__file__))
FIELDS = ["seq", "ms", "valid", "V1", "V2", "V3", "V4",
          "T1", "T2", "T3", "T4", "I", "STACK", "PACK", "LD"]
ROWS = 10
FIRST_ROW = 2
LAST_COL = "O"
UNO_PORT = 2002
XLSX_FILTER = "Calc MS Excel 2007 XML"


def parse_args():
    p = argparse.ArgumentParser(description="PCB data -> live LibreOffice Calc buffer")
    p.add_argument("--port", type=int, default=5011, help="receiver.py 가 넘겨주는 로컬 UDP 포트")
    p.add_argument("--dir", default=os.path.join(HERE, "buffer"))
    p.add_argument("--label", default="정상", help="데이터 구분 (파일 이름·시트 이름·표시 칸)")
    return p.parse_args()


def prop(name, value):
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


def make_template(path, label):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = label
    ws.append(FIELDS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    ws["Q1"], ws["Q2"], ws["Q3"], ws["Q4"] = "블록 번호", "채운 행", "마지막 저장", "구분"
    ws["R4"] = label
    for k in ("Q1", "Q2", "Q3", "Q4"):
        ws[k].font = Font(bold=True)
    ws["R1"], ws["R2"] = 0, 0
    for col, w in zip("ABCDEFGHIJKLMNO", [8, 9, 6] + [8] * 12):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["Q"].width = 12
    ws.column_dimensions["R"].width = 22
    ws.freeze_panes = "A2"
    wb.save(path)


def gui_env():
    env = dict(os.environ)
    env.setdefault("XDG_RUNTIME_DIR", "/run/user/%d" % os.getuid())
    env.setdefault("WAYLAND_DISPLAY", "wayland-0")
    env.setdefault("DBUS_SESSION_BUS_ADDRESS", "unix:path=%s/bus" % env["XDG_RUNTIME_DIR"])
    return env


class Calc:
    def __init__(self, path):
        self.path = path
        self.url = uno.systemPathToFileUrl(path)
        self.doc = self.sheet = None

    def _desktop(self):
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", local)
        ctx = resolver.resolve(f"uno:socket,host=127.0.0.1,port={UNO_PORT};urp;"
                               "StarOffice.ComponentContext")
        return ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)

    def connect(self):
        try:
            desktop = self._desktop()
        except Exception:
            subprocess.Popen(["soffice", "--calc", "--norestore", "--nologo", "--nodefault",
                              f"--accept=socket,host=127.0.0.1,port={UNO_PORT};urp;"],
                             env=gui_env(), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                             start_new_session=True)
            for _ in range(60):
                time.sleep(1)
                try:
                    desktop = self._desktop()
                    break
                except Exception:
                    pass
            else:
                raise RuntimeError("LibreOffice 에 연결하지 못했습니다")

        doc = None
        comps = desktop.getComponents().createEnumeration()
        while comps.hasMoreElements():
            c = comps.nextElement()
            if getattr(c, "URL", "") == self.url:
                doc = c
                break
        if doc is None:
            doc = desktop.loadComponentFromURL(self.url, "_blank", 0, ())
        self.doc = doc
        self.sheet = doc.Sheets.getByIndex(0)
        print(f"calc: {self.path} 연결됨", flush=True)

    def write_row(self, i, cols):
        r = FIRST_ROW + i
        self.sheet.getCellRangeByName(f"A{r}:{LAST_COL}{r}").setDataArray((tuple(cols),))
        self.sheet.getCellRangeByName("R2").setValue(i + 1)

    def set_block(self, n):
        self.sheet.getCellRangeByName("R1").setValue(n)

    def clear(self):
        self.sheet.getCellRangeByName(f"A{FIRST_ROW}:{LAST_COL}{FIRST_ROW + ROWS - 1}").clearContents(
            1 | 2 | 4 | 16)
        self.sheet.getCellRangeByName("R2").setValue(0)
        self.doc.setModified(False)

    def save_full(self, full_path, stamp):
        self.sheet.getCellRangeByName("R3").setString(stamp)
        tmp = full_path + ".tmp.xlsx"
        self.doc.storeToURL(uno.systemPathToFileUrl(tmp),
                            (prop("FilterName", XLSX_FILTER), prop("Overwrite", True)))
        os.replace(tmp, full_path)


def to_cell(i, s):
    try:
        return float(int(s)) if i < 2 else float(s)
    except ValueError:
        return s


def main():
    a = parse_args()
    os.makedirs(a.dir, exist_ok=True)
    path = os.path.join(a.dir, f"buffer({a.label}).xlsx")
    full_path = os.path.join(a.dir, f"buffer_full({a.label}).xlsx")
    if not os.path.exists(path):
        make_template(path, a.label)

    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    sock.bind(("127.0.0.1", a.port))

    calc = Calc(path)
    calc.connect()
    calc.clear()
    block, filled, last_seq = 1, 0, None
    calc.set_block(block)
    print(f"calc_buffer: UDP 127.0.0.1:{a.port} -> {path} (10행마다 {full_path})", flush=True)

    while True:
        data, _ = sock.recvfrom(65535)
        cols = data.decode("utf-8", "replace").strip().split(",")
        if len(cols) != len(FIELDS):
            continue
        try:
            seq = int(cols[0])
        except ValueError:
            continue
        try:
            if last_seq is not None and seq <= last_seq and filled:
                print(f"seq 되돌아감 {last_seq}->{seq}: 채우던 {filled}행 버리고 새로 시작", flush=True)
                calc.clear()
                filled = 0
            last_seq = seq
            calc.write_row(filled, [to_cell(i, c) for i, c in enumerate(cols)])
            filled += 1
            if filled == ROWS:
                stamp = time.strftime("%Y-%m-%d %H:%M:%S")
                calc.save_full(full_path, stamp)
                block += 1
                calc.clear()
                calc.set_block(block)
                calc.doc.setModified(False)
                filled = 0
        except Exception as e:
            print(f"calc 오류 ({e.__class__.__name__}: {e}), 다시 연결합니다", flush=True)
            time.sleep(1)
            try:
                calc.connect()
                calc.clear()
                calc.set_block(block)
            except Exception as e2:
                print(f"재연결 실패: {e2}", flush=True)
            filled = 0


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        pass
