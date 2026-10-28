#!/usr/bin/env python3
import argparse
import os
import socket
import time

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
FIELDS = ["seq", "ms", "valid", "V1", "V2", "V3", "V4",
          "T1", "T2", "T3", "T4", "I", "STACK", "PACK", "LD"]
IDX = {f: i for i, f in enumerate(FIELDS)}


def parse_args():
    p = argparse.ArgumentParser(description="Excel(정상) blocks -> AE monitor")
    p.add_argument("--label", default="정상")
    p.add_argument("--dir", default=os.path.join(HERE, "buffer"))
    p.add_argument("--monitor", default="127.0.0.1:5005", help="AE 모니터 입력 버퍼 UDP 주소")
    p.add_argument("--src", default="PCB", help="모니터에 표시할 출처 (PCB = PCB 실측)")
    p.add_argument("--poll", type=float, default=0.1, help="파일 확인 주기(초)")
    return p.parse_args()


def read_block(path):
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=11, max_col=len(FIELDS), values_only=True):
        if r[0] is None:
            continue
        rows.append(list(r))
    return rows


def to_line(src, r):
    v = lambda k: float(r[IDX[k]])
    vals = [v("ms") / 1000.0, v("V1"), v("V2"), v("V3"), v("V4"), v("I"),
            v("T1"), v("T2"), v("T3"), v("T4"), 0, 0, 0, 0]
    return f"{src},{int(r[IDX['seq']])}," + ",".join(f"{x:.6g}" for x in vals)


def main():
    a = parse_args()
    path = os.path.join(a.dir, f"buffer_full({a.label}).xlsx")
    host, port = a.monitor.rsplit(":", 1)
    dst = (host, int(port))
    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    print(f"excel_to_monitor: {path} -> 모니터 UDP {dst[0]}:{dst[1]} (SRC={a.src})", flush=True)

    last_mtime, last_seq, n_blocks = None, None, 0
    while True:
        time.sleep(a.poll)
        try:
            mtime = os.stat(path).st_mtime_ns
        except FileNotFoundError:
            continue
        if mtime == last_mtime:
            continue
        last_mtime = mtime
        try:
            rows = read_block(path)
        except Exception as e:
            print(f"읽기 실패 ({e.__class__.__name__}), 다음 블록에서 다시", flush=True)
            last_mtime = None
            continue
        if not rows:
            continue

        first = int(rows[0][0])
        if last_seq is not None and first <= last_seq - 10:
            print(f"seq 되돌아감 {last_seq}->{first}: 새 스트림", flush=True)
            last_seq = None
        new = [r for r in rows if last_seq is None or int(r[0]) > last_seq]
        if not new:
            continue
        if last_seq is not None and int(new[0][0]) > last_seq + 1:
            print(f"블록 사이 누락: seq {last_seq + 1}~{int(new[0][0]) - 1}", flush=True)
        try:
            payload = "\n".join(to_line(a.src, r) for r in new) + "\n"
        except (TypeError, ValueError) as e:
            print(f"값 변환 실패 ({e}), 블록 건너뜀", flush=True)
            continue
        sock.sendto(payload.encode(), dst)
        last_seq = int(new[-1][0])
        n_blocks += 1
        if n_blocks % 30 == 1:
            print(f"블록 {n_blocks}: seq {int(new[0][0])}~{last_seq} ({len(new)}행) 전송", flush=True)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        pass
